import sys
from docx import Document
from docx.oxml.ns import qn


doc = Document(sys.argv[1])


def walk_paragraphs():
    for i, p in enumerate(doc.paragraphs):
        yield f"body:{i}", p
    for ti, table in enumerate(doc.tables):
        seen = set()
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                for pi, p in enumerate(cell.paragraphs):
                    yield f"table:{ti}:{ri}:{ci}:{pi}", p


for loc, p in walk_paragraphs():
    instr = [e.text for e in p._p.iter(qn("w:instrText")) if e.text]
    if instr or "SEQ" in p.text or p.style.name == "Caption":
        all_text = "".join((e.text or "") for e in p._p.iter() if e.tag in {qn("w:t"), qn("w:instrText")})
        print(loc, "style=", p.style.name, "ptext=", repr(p.text), "all=", repr(all_text), "instr=", instr)
