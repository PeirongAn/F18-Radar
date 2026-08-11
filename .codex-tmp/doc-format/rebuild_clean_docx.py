from __future__ import annotations

import io
import sys
from copy import deepcopy

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def remove_all_children(element):
    for child in list(element):
        element.remove(child)


def sanitize_section_refs(root):
    for sectpr in root.iter(qn("w:sectPr")):
        for tag in ("w:headerReference", "w:footerReference"):
            for node in list(sectpr.findall(qn(tag))):
                sectpr.remove(node)


def map_image_relationships(src, dst):
    mapping = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype != RT.IMAGE:
            continue
        blob = rel.target_part.blob
        new_rid, _ = dst.part.get_or_add_image(io.BytesIO(blob))
        mapping[rid] = new_rid
    return mapping


def remap_relationship_attributes(root, mapping):
    for element in root.iter():
        for attr_name, value in list(element.attrib.items()):
            if value in mapping:
                element.set(attr_name, mapping[value])


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(begin)
    paragraph._p.append(begin_run)

    instr_run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    instr_run.append(instr)
    paragraph._p.append(instr_run)

    sep_run = OxmlElement("w:r")
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    sep_run.append(sep)
    paragraph._p.append(sep_run)

    result_run = OxmlElement("w:r")
    result = OxmlElement("w:t")
    result.text = "1"
    result_run.append(result)
    paragraph._p.append(result_run)

    end_run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)
    paragraph._p.append(end_run)


def rebuild_footers(doc):
    for section in doc.sections:
        section.footer.is_linked_to_previous = False
        footer = section.footer
        for p in footer.paragraphs:
            remove_all_children(p._p)
        p = footer.paragraphs[0]
        add_page_number(p)
        sectpr = section._sectPr
        pg_num = sectpr.find(qn("w:pgNumType"))
        if pg_num is None:
            pg_num = OxmlElement("w:pgNumType")
            sectpr.append(pg_num)
        pg_num.set(qn("w:start"), "1")


def main(src_path, dst_path):
    src = Document(src_path)
    dst = Document()

    dst._part._styles_part._element = deepcopy(src._part._styles_part._element)
    image_rids = map_image_relationships(src, dst)

    src_body = src._body._element
    dst_body = dst._body._element
    remove_all_children(dst_body)
    for child in src_body:
        copied = deepcopy(child)
        sanitize_section_refs(copied)
        remap_relationship_attributes(copied, image_rids)
        dst_body.append(copied)

    rebuild_footers(dst)
    settings = dst.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    dst.save(dst_path)
    print(dst_path)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
