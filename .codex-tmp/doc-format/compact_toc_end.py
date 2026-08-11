import sys

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


doc = Document(sys.argv[1])
paragraphs = doc.paragraphs
first_body_heading = next(
    i for i, p in enumerate(paragraphs)
    if p.style and p.style.name == "Heading 1"
)

section_break_index = next(
    i for i in range(first_body_heading - 1, -1, -1)
    if paragraphs[i]._p.pPr is not None and paragraphs[i]._p.pPr.sectPr is not None
)

# Remove purely empty spacer paragraphs between the TOC field and the section break.
for p in list(paragraphs[:section_break_index]):
    if p.text or p._p.pPr is not None and p._p.pPr.sectPr is not None:
        continue
    has_field = any(el.tag == qn("w:fldChar") for el in p._p.iter())
    if has_field:
        continue
    # Only remove spacers after the last TOC entry.
    if p._p.getprevious() is not None:
        prev_text = "".join(p._p.getprevious().itertext())
        if "PAGEREF" in prev_text or p._p.getnext() is paragraphs[section_break_index]._p:
            p._p.getparent().remove(p._p)

# Re-read after removal, then compact the invisible field-end and section-break paragraphs.
paragraphs = doc.paragraphs
first_body_heading = next(
    i for i, p in enumerate(paragraphs)
    if p.style and p.style.name == "Heading 1"
)
section_break_index = next(
    i for i in range(first_body_heading - 1, -1, -1)
    if paragraphs[i]._p.pPr is not None and paragraphs[i]._p.pPr.sectPr is not None
)

for i in range(max(0, section_break_index - 2), section_break_index + 1):
    p = paragraphs[i]
    if p.text:
        continue
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(1)
    ppr = p._p.get_or_add_pPr()
    outline = ppr.find(qn("w:outlineLvl"))
    if outline is not None:
        ppr.remove(outline)

doc.save(sys.argv[2])
print(sys.argv[2])
