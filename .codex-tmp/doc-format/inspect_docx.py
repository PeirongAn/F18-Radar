from __future__ import annotations

import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def emu_to_cm(value):
    return None if value is None else round(value / 360000, 3)


def pt(value):
    return None if value is None else round(value.pt, 2)


def paragraph_record(i, p):
    fmt = p.paragraph_format
    runs = []
    for r in p.runs:
        if not r.text:
            continue
        rpr = r._element.rPr
        east = None
        ascii_font = None
        if rpr is not None and rpr.rFonts is not None:
            east = rpr.rFonts.get(qn("w:eastAsia"))
            ascii_font = rpr.rFonts.get(qn("w:ascii"))
        runs.append(
            {
                "text": r.text[:120],
                "bold": r.bold,
                "italic": r.italic,
                "underline": bool(r.underline) if r.underline is not None else None,
                "size": pt(r.font.size),
                "font": r.font.name,
                "ascii_font": ascii_font,
                "east_asia_font": east,
                "color": str(r.font.color.rgb) if r.font.color and r.font.color.rgb else None,
            }
        )
    return {
        "index": i,
        "text": p.text[:300],
        "style": p.style.name if p.style else None,
        "alignment": int(p.alignment) if p.alignment is not None else None,
        "left_indent_cm": emu_to_cm(fmt.left_indent),
        "right_indent_cm": emu_to_cm(fmt.right_indent),
        "first_line_indent_cm": emu_to_cm(fmt.first_line_indent),
        "space_before_pt": pt(fmt.space_before),
        "space_after_pt": pt(fmt.space_after),
        "line_spacing": str(fmt.line_spacing) if fmt.line_spacing is not None else None,
        "keep_with_next": fmt.keep_with_next,
        "page_break_before": fmt.page_break_before,
        "runs": runs,
    }


def main(src, out):
    doc = Document(src)
    records = [paragraph_record(i, p) for i, p in enumerate(doc.paragraphs)]
    style_counts = Counter(r["style"] for r in records if r["text"].strip())
    direct = Counter()
    for r in records:
        for run in r["runs"]:
            key = (run["font"], run["east_asia_font"], run["size"], run["bold"], run["italic"], run["color"])
            direct[key] += 1

    sections = []
    for i, s in enumerate(doc.sections):
        sections.append(
            {
                "index": i,
                "start_type": int(s.start_type),
                "page_width_cm": emu_to_cm(s.page_width),
                "page_height_cm": emu_to_cm(s.page_height),
                "top_margin_cm": emu_to_cm(s.top_margin),
                "bottom_margin_cm": emu_to_cm(s.bottom_margin),
                "left_margin_cm": emu_to_cm(s.left_margin),
                "right_margin_cm": emu_to_cm(s.right_margin),
                "header_distance_cm": emu_to_cm(s.header_distance),
                "footer_distance_cm": emu_to_cm(s.footer_distance),
            }
        )

    tables = []
    for i, t in enumerate(doc.tables):
        tables.append({"index": i, "rows": len(t.rows), "cols": len(t.columns), "style": t.style.name if t.style else None})

    payload = {
        "paragraph_count": len(records),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "style_counts": style_counts,
        "direct_run_formats": [{"format": k, "count": v} for k, v in direct.most_common()],
        "sections": sections,
        "tables": tables,
        "paragraphs": records,
    }
    Path(out).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: payload[k] for k in ("paragraph_count", "table_count", "inline_shape_count", "style_counts", "sections")}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
