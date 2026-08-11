from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


CAPTIONS = {
    10: "图 1 战场态势特征分析层次结构",
    32: "图 2 任务需求特征分析",
    55: "图 3 智能体自主能力分析",
    78: "表 1 交互通道分析",
    81: "图 4 交互通道优先级策略",
    83: "表 2 交互通道降级策略",
    87: "表 3 人机协同模式划分",
    90: "表 4 模式转换触发条件矩阵",
    96: "表 5 信息反馈设计原则",
    100: "图 5 AI决策可解释层级示例",
    124: "表 6 远程态势层表征参数",
    127: "表 7 中层态势层表征参数",
    130: "表 8 近程态势层表征参数",
    137: "表 9 任务执行环节关键数据",
    141: "表 10 任务绩效体系说明",
    143: "表 11 平台控制类任务特征参数与评价指标",
    145: "表 12 传感器交互类任务特征参数与评价指标",
    147: "表 13 威胁排序类任务特征参数与评价指标",
    149: "表 14 武器发射类任务特征参数与评价指标",
    151: "图 6 四类任务的研究变量与观测指标框架",
    156: "表 15 个体差异与环境约束参数设计表",
    163: "图 7 三种任务模式下的人机功能分配关系",
    203: "图 8 典型空战任务仿真与交互等效试验平台架构",
    209: "图 9 平台数据流转机制说明",
}


def set_run_font(run, size=10.5, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "宋体")
    for highlight in list(rpr.findall(qn("w:highlight"))):
        rpr.remove(highlight)


def clear_paragraph_content(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def rebuild_toc_field(p):
    clear_paragraph_content(p)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(begin)
    p._p.append(begin_run)

    instr_run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    instr_run.append(instr)
    p._p.append(instr_run)

    sep_run = OxmlElement("w:r")
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    sep_run.append(sep)
    p._p.append(sep_run)

    placeholder_run = OxmlElement("w:r")
    placeholder = OxmlElement("w:t")
    placeholder.text = '目录将在打开文档时自动生成，请右键选择“更新域”刷新。'
    placeholder_run.append(placeholder)
    p._p.append(placeholder_run)

    end_run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)
    p._p.append(end_run)


def rebuild_caption(p, text):
    clear_paragraph_content(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=10.5, bold=False)


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if fill is None:
        if shd is not None:
            tcpr.remove(shd)
        return
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=55, start=65, bottom=55, end=65):
    tcpr = cell._tc.get_or_add_tcPr()
    tc_mar = tcpr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = qn(f"w:{edge}")
        node = tc_mar.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_problem_table(table):
    table.autofit = False
    seen = set()
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            key = cell._tc
            if key in seen:
                continue
            seen.add(key)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, "D9E5F0" if row_index == 0 else None)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_run_font(run, size=9.5, bold=(row_index == 0))

    trpr = table.rows[0]._tr.get_or_add_trPr()
    if trpr.find(qn("w:tblHeader")) is None:
        trpr.append(OxmlElement("w:tblHeader"))


def set_update_fields(doc):
    settings = doc.settings._element
    for node in list(settings.findall(qn("w:showFieldCodes"))):
        settings.remove(node)
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main(src, dst):
    doc = Document(src)
    rebuild_toc_field(doc.paragraphs[0])
    for index, text in CAPTIONS.items():
        rebuild_caption(doc.paragraphs[index], text)
    format_problem_table(doc.tables[14])
    set_update_fields(doc)
    Path(dst).parent.mkdir(parents=True, exist_ok=True)
    doc.save(dst)
    print(dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
