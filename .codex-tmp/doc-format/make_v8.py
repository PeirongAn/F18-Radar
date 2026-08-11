from __future__ import annotations

import sys
from copy import deepcopy

from docx import Document
from docx.oxml.ns import qn


RENAMES = {
    "2.3.1平台自主控制能力": "2.3.1平台控制子任务的自主控制能力",
    "2.3.2目标识别与锁定能力": "2.3.2传感器交互子任务的目标识别与锁定能力",
    "2.3.3威胁评估与排序能力": "2.3.3威胁排序子任务的威胁评估与排序能力",
    "2.3.4武器发射决策能力": "2.3.4武器发射子任务的发射决策能力",
}


INSERTIONS = [
    (
        "平台控制子任务以飞行平台的航路跟踪和姿态保持为研究载体",
        "2.5.1平台控制子任务场景概述",
    ),
    (
        "传感器交互子任务以雷达操作为研究载体",
        "2.5.2传感器交互子任务场景概述",
    ),
    (
        "威胁排序子任务以多目标条件下的首要威胁识别为研究载体",
        "2.5.3威胁排序子任务场景概述",
    ),
    (
        "武器发射子任务以发射时机判断和发射执行为研究载体",
        "2.5.4武器发射子任务场景概述",
    ),
    (
        "平台控制任务的评价应围绕任务结果、控制参与和协同过程构建",
        "3.3.1平台控制任务绩效评价指标",
    ),
    (
        "传感器任务可从阶段时长、操作正确性、操作主体和认知状态四方面评价",
        "3.3.2传感器交互任务绩效评价指标",
    ),
    (
        "评价指标应与最高威胁识别范式保持一致",
        "3.3.3威胁排序任务绩效评价指标",
    ),
    (
        "武器发射任务可采用发射次数、发射成功次数、发射成功率",
        "3.3.4武器发射任务绩效评价指标",
    ),
]


def copy_run_format(src_run, dst_run):
    src_rpr = src_run._r.rPr
    if src_rpr is None:
        return
    dst_rpr = dst_run._r.get_or_add_rPr()
    parent = dst_rpr.getparent()
    parent.replace(dst_rpr, deepcopy(src_rpr))


def replace_text_preserve_format(paragraph, text):
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = text
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
    else:
        paragraph.add_run(text)


def main(src_path, dst_path):
    doc = Document(src_path)
    heading_template = next(
        p for p in doc.paragraphs if p.style and p.style.name == "Heading 3" and p.runs
    )

    renamed = set()
    for p in doc.paragraphs:
        if not p.style or p.style.name != "Heading 3":
            continue
        old = p.text.strip()
        if old in RENAMES:
            replace_text_preserve_format(p, RENAMES[old])
            renamed.add(old)
    missing_renames = set(RENAMES) - renamed
    if missing_renames:
        raise RuntimeError(f"Missing headings to rename: {sorted(missing_renames)}")

    inserted = []
    for prefix, heading_text in INSERTIONS:
        existing = [p for p in doc.paragraphs if p.text.strip() == heading_text]
        if existing:
            raise RuntimeError(f"Heading already exists: {heading_text}")
        targets = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
        if len(targets) != 1:
            raise RuntimeError(f"Expected one target for {heading_text}, found {len(targets)}")
        new_p = targets[0].insert_paragraph_before(heading_text, style="Heading 3")
        new_p.paragraph_format.keep_with_next = True
        if new_p.runs:
            copy_run_format(heading_template.runs[0], new_p.runs[0])
        inserted.append(heading_text)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        from docx.oxml import OxmlElement

        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    doc.core_properties.title = "专题1-1典型智能KZ任务场景特征分析研究报告V8"
    doc.save(dst_path)
    print(f"renamed={len(renamed)} inserted={len(inserted)} output={dst_path}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
