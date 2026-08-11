from __future__ import annotations

import sys
from docx import Document
from docx.oxml.ns import qn


def attr(el, tag, attr):
    found = el.find(qn(tag))
    return None if found is None else found.get(qn(attr))


doc = Document(sys.argv[1])
for ti in [12, 13, 14]:
    t = doc.tables[ti]
    print(f"TABLE {ti} rows={len(t.rows)} cols={len(t.columns)} style={t.style.name if t.style else None}")
    print("grid", [x.get(qn("w:w")) for x in t._tbl.tblGrid.gridCol_lst])
    for ri in range(min(2, len(t.rows))):
        vals = []
        for ci, cell in enumerate(t.rows[ri].cells):
            tcpr = cell._tc.tcPr
            shade = attr(tcpr, "w:shd", "w:fill")
            width = tcpr.tcW.w if tcpr.tcW is not None else None
            p = cell.paragraphs[0]
            runs = []
            for run in p.runs:
                highlight = None
                if run._element.rPr is not None:
                    hl = run._element.rPr.find(qn("w:highlight"))
                    highlight = None if hl is None else hl.get(qn("w:val"))
                runs.append((run.text[:24], run.font.name, run.font.size.pt if run.font.size else None, run.bold, highlight))
            vals.append({"c": ci, "w": width, "shade": shade, "align": int(p.alignment) if p.alignment is not None else None, "runs": runs})
        print("row", ri, vals)
