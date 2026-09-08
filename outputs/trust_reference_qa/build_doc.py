from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
out=Path(r'D:\codes\F18-Radar\outputs\人机信任调控因素及试验设计参考.docx')
doc=Document()
s=doc.sections[0]
s.page_width=Cm(21); s.page_height=Cm(29.7)
s.top_margin=Cm(2.4); s.bottom_margin=Cm(2.4); s.left_margin=Cm(2.6); s.right_margin=Cm(2.6)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[name]; st.font.name='宋体'; st.font.size=Pt(10.5); st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
 pf=st.paragraph_format; pf.space_before=Pt(0); pf.space_after=Pt(0); pf.line_spacing=1.35
 if name=='Normal': pf.first_line_indent=Pt(21)
 else: pf.first_line_indent=Pt(0); st.font.bold=True; pf.keep_with_next=True
for name in ['Heading 1','Heading 2']:
 doc.styles[name].paragraph_format.space_before=Pt(6)

def p(t):
 x=doc.add_paragraph(t); x.paragraph_format.widow_control=True
 x.paragraph_format.first_line_indent=Pt(21)
 x._p.get_or_add_pPr().get_or_add_ind().set(qn('w:firstLineChars'),'200')
 return x

def h(t): doc.add_paragraph(t,'Heading 1')
def sub(t): doc.add_paragraph(t,'Heading 2')
t=doc.add_paragraph('人机信任调控因素及试验设计参考','Title'); t.alignment=WD_ALIGN_PARAGRAPH.CENTER; t.paragraph_format.space_after=Pt(10)
p('本材料结合传感器任务和威胁排序任务中的现有实现，归纳可供类似试验参考的调控因素、触发依据和评价口径。调控目标是帮助操作者合理采纳或拒绝人工智能（AI）建议，减少欠信任与过信任。现有功能可提炼为五类因素，具体效果及适用条件仍需通过试验验证。')
h('一 调控因素及其作用方式')
sub('1 AI 能力信息呈现')
p('向操作者展示 AI 能力等级对应的准确率参考值及波动范围，为判断系统能力提供信息。可设置不展示、展示单一参考值、展示参考值及波动范围等试验条件，考察能力信息对建议采纳和信任判断的影响。当前曲线按 AI 等级概率范围生成，应明确标注为能力参考，不能作为实测历史正确率。')
sub('2 决策信息透明度')
p('在推荐结果之外展示目标观测信息。传感器任务包括方位、距离、速度、航向及数据时间等信息，并提供候选目标列表；威胁排序任务包括威胁类别、目标特征、距离及数据时间等信息。可比较仅展示推荐结果与同时展示相关观测信息的条件。当前实现属于信息辅助，不等同于完整的 AI 推理过程解释。')
sub('3 注意引导')
p('进入信任支持状态后，对 AI 推荐区域进行高亮，引导操作者关注推荐信息；聚焦推荐目标或开始人工复核后结束高亮。可将是否高亮、触发时机和结束条件作为试验变量，考察推荐信息的注意分配和查看行为。')
sub('4 人工复核支持')
p('通过专用按键查看 AI 推荐目标的观测信息，记录有效复核次数、持续时间及无有效推荐时的复核操作。可比较是否提供复核入口或不同复核方式的影响。当前复核用于辅助判断，不能直接等同于操作者已经理解信息，也不能仅凭复核次数判断信任是否适当。')
sub('5 人机分歧对比')
p('当人工选择与 AI 推荐不一致时，并列展示两个目标的观测信息，支持操作者比较差异。可设置是否提供对比、对比字段范围等条件，考察分歧后的选择调整和最终决策。人机选择不一致只表示发生分歧，不应直接判定为欠信任。')
p('以上试验变量由现有功能提炼而来，供试验设计选取；不代表当前系统已具备各因素的独立开关，也不代表已完成单因素效果验证。')
doc.add_page_break()
h('二 试验条件与调控触发')
p('试验条件建议至少记录任务类型、任务难度、AI 能力等级和操作模式。操作模式可区分作业员独立操作、智能体自主操作和人机协同操作。涉及 AI 建议采纳或拒绝的信任行为分析，应限定在操作者有机会作出相关选择的条件下；练习阶段与正式试验阶段应分别记录。')
p('当前调控依据同一被试、同一任务类型、同一任务难度和同一 AI 等级的历史结果，统计信任适当、欠信任和过信任情况。历史中出现欠信任或过信任后进入信任支持状态，两类偏差目前采用共同的支持模式。该规则是现有实现口径，不能表述为已经验证的最优触发阈值，也不宜描述为两套已独立验证的定向干预策略。')
p('对方设计试验时，可先比较基础界面与调控界面的整体效果，再按研究目的拆分能力信息、透明度、注意引导、复核支持和分歧对比等因素。建议明确每项因素的启用条件、呈现内容和结束条件，避免多个因素同时变化后难以判断效果来源。')
h('三 评价指标与判定口径')
sub('1 信任行为结果')
p('以 AI 推荐、人工最终选择和任务真值为判定依据：AI 正确且人工采纳，记为信任适当；AI 正确但人工拒绝，记为欠信任；AI 错误且人工采纳，记为过信任；AI 错误而人工拒绝，记为信任适当。此处“适当”评价的是对 AI 建议的依赖行为。人工拒绝错误建议后仍可能选错，因此必须另行统计任务正确率。')
p('分别计算信任适当率、欠信任率和过信任率，分母为能够完成上述判定的有效试次。AI 推荐、人工最终选择或任务真值缺失时，不宜强行归类，应单独记录缺失原因。')
sub('2 任务绩效与过程行为')
p('任务绩效可记录最终选择正确率及决策耗时；决策耗时应统一起止事件，例如从 AI 推荐呈现到人工最终确认。过程行为可记录人工复核次数和时长、观测信息查看时长、人机分歧对比的出现与持续时间、选择变化以及相关区域的眼动行为。过程指标用于解释调控如何影响决策，不直接替代信任结果判定。')
sub('3 主观评价')
p('可结合研究目的配置任务组结束后的信任问卷，将主观信任与行为信任分别分析，并统一施测时点。问卷呈现与问卷提交应分别记录，避免将已触发但未提交的问卷计入有效样本。')
h('四 试验记录与使用边界')
p('建议按试次关联被试标识、任务组及试次标识、任务条件、调控因素及参数、实际触发记录、AI 推荐、人工最终选择、任务真值和关键事件时间。问卷及眼动数据通过相应标识与任务记录关联，以支持比较调控前后的行为和绩效。')
p('本材料可作为类似试验的因素选取和记录口径参考。迁移时需结合对方任务中的可观测信息、真值定义和人工决策权限细化方案；调控是否有效、适用于何种被试与任务，以及各因素是否存在交互作用，应以相应试验结果为依据。')
doc.core_properties.title='人机信任调控因素及试验设计参考'; doc.core_properties.author=''; doc.core_properties.subject='调控因素与试验设计参考'; doc.core_properties.keywords=''
for st in doc.styles:
 for border in list(st.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
 if st.type == 1:
  st.font.name='宋体'; st.font.size=Pt(10.5); st.font.color.rgb=RGBColor(0,0,0)
  rf=st.element.get_or_add_rPr().rFonts
  for attr in list(rf.attrib):
   if 'Theme' in attr: del rf.attrib[attr]
  for a in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+a),'宋体')
for border in list(doc.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
doc.save(out)
print(out)


