from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "/Users/anpeirong/code/radar/F18-Radar/outputs/传感器任务与威胁排序_三项版.docx"
FIG1 = "/Users/anpeirong/code/radar/F18-Radar/outputs/docx_media/figure_1.png"
FIG3 = "/Users/anpeirong/code/radar/F18-Radar/outputs/docx_media/figure_3.png"


def set_east_asia_font(run, font_name="宋体"):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.1):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                set_paragraph_spacing(paragraph, after=2, line=1.1)
                for run in paragraph.runs:
                    set_east_asia_font(run)


def add_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    set_paragraph_spacing(p, before=16 if level == 1 else 10, after=6, line=1.1)
    for run in p.runs:
        set_east_asia_font(run, "黑体")
        run.font.color.rgb = RGBColor(46, 116, 181 if level < 3 else 120)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=6, line=1.1)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        set_east_asia_font(r1, "宋体")
        r2 = p.add_run(text[len(bold_prefix):])
        set_east_asia_font(r2, "宋体")
    else:
        run = p.add_run(text)
        set_east_asia_font(run, "宋体")
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_paragraph_spacing(p, after=4, line=1.167)
        run = p.add_run(item)
        set_east_asia_font(run, "宋体")


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_paragraph_spacing(p, after=4, line=1.167)
        run = p.add_run(item)
        set_east_asia_font(run, "宋体")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for idx, h in enumerate(headers):
        header_cells[idx].text = h
        set_cell_shading(header_cells[idx], "F2F4F7")
        for p in header_cells[idx].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                set_east_asia_font(r, "黑体")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            for p in cells[idx].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    set_east_asia_font(r, "宋体")
    set_table_width(table, widths)
    doc.add_paragraph()
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=2, after=8, line=1.1)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_east_asia_font(r, "宋体")
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(85, 85, 85)
    return p


def add_picture(doc, path, caption):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=4, after=2, line=1.0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Inches(6.35))
    add_caption(doc, caption)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
styles["Normal"].font.size = Pt(11)
for name in ("Heading 1", "Heading 2", "Heading 3"):
    styles[name].font.name = "Calibri"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

title = doc.add_paragraph()
set_paragraph_spacing(title, after=4, line=1.1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("传感器任务与威胁排序信任调控方案")
set_east_asia_font(run, "黑体")
run.font.size = Pt(18)
run.bold = True
run.font.color.rgb = RGBColor(11, 37, 69)

subtitle = doc.add_paragraph()
set_paragraph_spacing(subtitle, after=12, line=1.1)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("三项版：信任状态判断、欠信任调控、过信任调控")
set_east_asia_font(r, "宋体")
r.font.size = Pt(10.5)
r.font.color.rgb = RGBColor(85, 85, 85)

add_heading(doc, "1. 信任状态如何判断", 1)
add_body(
    doc,
    "信任状态识别的目标，是在不打断操作员任务的前提下，依据可观测交互行为判断操作员当前属于欠信任、校准信任还是过信任，并据此决定是否触发界面调控。判断对象是“操作员如何使用AI建议”，不是直接判断AI建议本身是否正确。",
)
add_table(
    doc,
    ["状态", "典型表现", "判定含义"],
    [
        ["欠信任", "持续拒绝或长时间不接受高置信AI建议；频繁接管、反复查看同一证据仍不确认。", "操作员对AI建议的信任低于当前AI能力和数据质量所支持的合理水平。"],
        ["校准信任", "会查看关键证据；高置信、数据质量好的建议可较快确认；低置信或高风险场景会复核。", "操作员信任水平与AI能力、数据质量和任务风险基本匹配。"],
        ["过信任", "连续接受AI建议；低置信、IFF未确认或候选接近时仍不查看证据、不执行复核。", "操作员对AI建议的依赖超过当前证据和风险条件所允许的水平。"],
    ],
    [1.1, 2.8, 2.6],
)
add_body(doc, "识别指标在最近N次交互或固定时间窗内统计，单次异常不直接改变状态。建议采用滑动窗口、持续阈值和回滞机制，避免状态在边界反复抖动。")
add_table(
    doc,
    ["识别维度", "观测指标", "状态指向"],
    [
        ["采纳行为", "AI建议接受率、拒绝次数、接管次数。", "接受率持续偏低指向欠信任；接受率过高且缺少复核指向过信任。"],
        ["验证行为", "解释面板查看率、证据展开次数、人工复核动作。", "反复查看仍不确认指向欠信任；几乎不查看证据即确认指向过信任。"],
        ["决策时延", "高置信建议确认时延、低置信建议确认时延。", "高置信仍长时间不确认指向欠信任；低置信也秒级确认指向过信任。"],
        ["一致性", "最终选择与AI建议一致率、人工修正频率。", "一致率持续偏低、频繁改选指向欠信任；一致率近乎100%且从不修正指向过信任。"],
    ],
    [1.35, 2.45, 2.7],
)
add_body(doc, "任务风险信号需要单独记录，不直接等同于信任状态。AI置信度低、候选目标概率差小、IFF未开启、目标聚集、航迹交叉、信号干扰或数据延迟，说明的是当前建议风险较高；它们应与操作员信任状态联合决定界面反应。")
add_numbered(
    doc,
    [
        "采集：记录每次AI建议及操作员响应，包括接受、拒绝、接管、查看证据和确认时延。",
        "统计：在滑动窗口内计算采纳、验证、时延和一致性指标。",
        "判定：任一维度越过阈值并持续存在时，输出对应信任状态；多维度冲突时优先参考验证行为和低风险/高风险场景下的一致性。",
        "触发：将信任状态与任务风险共同输入调控规则，决定是否展开解释、显示不确定性或增加复核门槛。",
    ],
)

add_heading(doc, "2. 欠信任调控：选用AI推荐目标透明度叠加", 1)
add_body(
    doc,
    "选用理由：该方案最适合作为欠信任调控的基础方案。它不强迫操作员接受AI，而是把AI为什么推荐某个目标、推荐把握有多大、哪些环节仍需人工确认直接呈现在原有传感器任务界面中，使操作员能够快速验证AI建议。",
)
add_body(doc, "触发条件：当操作员连续拒绝AI推荐、长时间不接受高置信建议、频繁接管重新选择目标，或反复查看日志但仍不确认时，自动展开该调控。")
add_body(
    doc,
    "调控目标：将AI推荐从“单一结论”转化为“可定位、可解释、可复核的协同建议”。对欠信任操作员而言，界面不应只强调“请相信AI”，而应提供足够证据，使其能在短时间内确认AI是否有合理依据。",
)
add_table(
    doc,
    ["设计项", "呈现内容", "呈现方式"],
    [
        ["推荐目标标识", "AI推荐目标编号、敌方概率、锁定置信度。", "在目标外使用蓝色虚线框，旁侧显示“AI推荐T-07 / 置信度87%”。"],
        ["依据摘要", "航向接近、速度变化异常、轨迹稳定性、IFF状态。", "右侧日志下方显示“AI目标锁定解释”卡片，默认展开前三条依据。"],
        ["不确定性保留", "IFF未确认、候选目标仍存在、数据质量限制。", "使用黄色标签标明“待人工确认”，避免把AI建议表达成最终结论。"],
        ["操作建议", "接受锁定、开启IFF、人工复核。", "在解释卡片底部给出下一步动作，保留人工接管入口。"],
    ],
    [1.3, 2.55, 2.65],
)
add_picture(doc, FIG1, "图1 欠信任调控：AI推荐目标透明度叠加示意")
add_heading(doc, "2.1 截图界面说明", 2)
add_body(doc, "截图中，雷达B显区域使用蓝色虚线框标出AI推荐目标T-07，并在目标旁显示“AI推荐T-07 / 置信度87%”。蓝色表示该目标来自AI推荐，虚线表示仍需人工确认，避免操作员误认为系统已经完成最终锁定。")
add_body(doc, "右侧“AI目标锁定解释”卡片集中展示敌方概率、航向接近、速度变化异常、IFF状态和推荐动作。该卡片承担解释功能，帮助操作员理解AI为何选择T-07，而不是让其在通信日志中自行推断。")
add_body(doc, "黄色虚线框标出备选目标，说明AI并非只给出孤立答案，而是在多个候选中进行了比较。对欠信任操作员来说，备选目标提示可以降低“AI是否遗漏其他目标”的疑虑。")
add_heading(doc, "2.2 交互流程", 2)
add_table(
    doc,
    ["步骤", "系统/操作员动作", "界面反馈", "调控作用"],
    [
        ["1", "系统检测到操作员连续拒绝AI推荐，或高置信建议长时间未确认。", "触发欠信任状态，右侧解释卡片从折叠变为展开。", "避免继续只给出结论，转为提供证据。"],
        ["2", "AI完成目标分析并推荐T-07。", "雷达中T-07出现蓝色虚线框，旁侧显示置信度87%。", "让操作员快速定位AI关注目标。"],
        ["3", "操作员查看目标或悬停推荐框。", "右侧解释卡片高亮敌方概率、航向接近、速度变化异常等依据。", "把推荐结果转化为可验证证据链。"],
        ["4", "操作员对推荐目标仍有疑问。", "界面保留黄色备选目标提示，可进一步查看候选差异。", "说明AI已比较其他可能目标，减少黑箱感。"],
        ["5", "操作员选择“接受锁定”或“开启IFF”。", "系统记录确认动作和所依据证据；确认后解释卡片自动折叠。", "恢复任务界面简洁度，并为后续状态识别提供日志。"],
    ],
    [0.55, 2.25, 2.25, 1.45],
)
add_heading(doc, "2.3 交互细节要求", 2)
add_body(doc, "界面逻辑：雷达B显区域只承担轻量标注，帮助操作员定位AI关注的目标；右侧解释面板承担证据说明，帮助操作员理解AI推荐的依据。该布局能在不遮挡主任务区域的前提下，提高AI建议的可理解性和可验证性。")
add_bullets(
    doc,
    [
        "蓝色表示“AI推荐”，不表示最终敌我确认；虚线表示结果仍处于待确认状态。",
        "高置信度可使用绿色条形显示，但必须同时显示数据质量和IFF状态，避免造成误导。",
        "当操作员接受推荐后，解释面板可自动折叠，避免长期占用通信日志区域。",
        "解释卡片默认显示3到5条关键依据，避免把通信日志区域变成大段说明文字。",
        "所有查看证据、开启IFF、接受锁定和人工接管动作都应写入信任调控事件日志，用于后续判断调控是否有效。",
    ],
)

add_heading(doc, "3. 过信任调控：选用低置信度与不确定性显性化", 1)
add_body(
    doc,
    "选用理由：该方案直接针对过信任的主要风险，即操作员把AI建议当作最终命令，在低置信、候选接近或IFF未确认时仍直接接受。通过降低确定性表达和增加复核门槛，可以把操作员从“自动接受”拉回到“监督判断”。",
)
add_body(doc, "触发条件：当AI置信度低于阈值、候选目标概率差小于设定值、目标聚集或航迹交叉、IFF尚未确认，同时操作员连续接受AI建议且未查看证据时，触发该调控。")
add_body(
    doc,
    "调控目标：降低低置信AI建议的确定性表达，并在关键风险条件下增加人工复核门槛。过信任调控不是让操作员不信任AI，而是提醒其当前AI输出仍是建议，不能替代人工确认。",
)
add_table(
    doc,
    ["设计项", "呈现内容", "呈现方式"],
    [
        ["低置信标识", "置信度低于70%、相似候选存在。", "推荐框由蓝色改为橙色虚线框，旁侧显示“置信度62%”。"],
        ["不确定性来源", "目标聚集、航迹交叉、IFF未确认、信号干扰。", "右侧弹出橙色“需要人工复核”卡片，避免使用绿色确认语义。"],
        ["反证提示", "其他目标敌方概率接近，或第二候选具有局部优势。", "在复核卡片中提示“T-07敌方概率接近，建议比较后再锁定”。"],
        ["强制复核", "接受锁定前必须查看证据或人工复核。", "低于阈值时禁用“一键接受”，显示“人工复核 / 查看证据”按钮。"],
    ],
    [1.3, 2.55, 2.65],
)
add_picture(doc, FIG3, "图2 过信任调控：低置信度与人工复核提示示意")
add_heading(doc, "3.1 截图界面说明", 2)
add_body(doc, "截图中，AI锁定目标T-11使用橙色虚线框标出，并显示“置信度62%”。橙色表示当前建议存在风险，虚线表示仍处于待复核状态；这会削弱“AI已给出最终答案”的暗示。")
add_body(doc, "雷达画面上方另有黄色“相似候选”目标，提示当前场景并非只有T-11符合敌方特征。该提示用于打断操作员对单一AI答案的惯性接受，促使其比较T-11与T-07等相近候选。")
add_body(doc, "右侧橙色卡片以“需要人工复核”为标题，列出当前AI建议置信度、不确定性来源、反证提示和阈值规则，并提供“人工复核”“查看证据”两个按钮。该卡片是过信任调控的核心组件。")
add_heading(doc, "3.2 交互流程", 2)
add_table(
    doc,
    ["步骤", "系统/操作员动作", "界面反馈", "调控作用"],
    [
        ["1", "系统检测到操作员连续直接接受AI建议，且很少查看证据。", "标记为过信任倾向，后续低置信建议进入复核模式。", "提前识别盲目依赖风险。"],
        ["2", "AI在目标聚集、航迹交叉、IFF未确认条件下给出T-11建议。", "推荐框改为橙色虚线，并显示置信度62%。", "降低确定性表达，提示当前建议不可直接执行。"],
        ["3", "系统发现T-07敌方概率接近。", "在雷达中标出相似候选，并在复核卡片中显示反证提示。", "迫使操作员比较候选，而不是只接受第一答案。"],
        ["4", "操作员尝试直接接受锁定。", "一键接受被禁用或转为二次确认，界面要求先选择“人工复核”或“查看证据”。", "把文字提示转化为实际交互约束。"],
        ["5", "操作员查看证据或执行人工复核。", "系统展开候选证据、IFF状态和不确定性来源；复核完成后恢复确认入口。", "确保接受动作建立在证据检查之后。"],
        ["6", "操作员最终确认或改选目标。", "系统记录是否查看证据、是否复核、最终选择与AI建议是否一致。", "为后续信任状态再判断提供依据。"],
    ],
    [0.55, 2.25, 2.25, 1.45],
)
add_heading(doc, "3.3 交互细节要求", 2)
add_body(doc, "界面逻辑：过信任调控不是否定AI能力，而是在证据不足时改变AI建议的语义强度。橙色、虚线框、相似候选提示和复核卡片共同表达“当前建议可参考，但不能直接作为最终结论”。")
add_bullets(
    doc,
    [
        "低置信或数据质量差时，不显示“已确认”“可直接锁定”等确定性强的语言。",
        "复核提示必须说明原因，例如置信度低于阈值、IFF未确认或航迹交叉，而不是只弹出笼统告警。",
        "对过信任状态，文字提醒不足以完成调控，应通过禁用一键接受、要求查看证据等交互流程形成约束。",
        "橙色提示用于需要注意和复核的状态，红色只保留给系统判定明显不可接受或严重错误的情况。",
        "复核完成后，系统可以恢复确认入口，但仍应保留人工接管选项，避免操作员被迫接受AI建议。",
    ],
)

add_heading(doc, "4. 简化后的调控规则", 1)
add_table(
    doc,
    ["状态/条件", "界面反应", "目的"],
    [
        ["欠信任 + 高置信/数据质量好", "展开AI推荐目标透明度叠加，显示目标、置信度和主要依据。", "降低黑箱感，恢复合理信任。"],
        ["过信任 + 低置信/高风险", "显示低置信与不确定性提示，增加人工复核门槛。", "防止盲目接受，恢复监督角色。"],
        ["正常信任 + 风险一般", "保持默认AI建议样式，只保留置信度和人工接管入口。", "避免过度提示干扰任务。"],
    ],
    [1.7, 2.75, 2.05],
)
add_body(doc, "本版文档仅保留一个欠信任调控方案和一个过信任调控方案。后续如需扩展，可再加入候选目标对比、能力边界提示或威胁排序专用方案。")

doc.save(OUT)
print(OUT)
