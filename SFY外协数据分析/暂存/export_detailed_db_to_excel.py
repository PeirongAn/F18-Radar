import sqlite3
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

def export_detailed_db_to_excel(db_path, output_file):
    """将数据库详细信息和示例数据导出到Excel文件"""
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    
    # 创建Excel工作簿
    wb = Workbook()
    
    # 删除默认工作表
    wb.remove(wb.active)
    
    # 获取所有表名
    cursor.execute("SELECT name FROM sqlite_master WHERE type='table';")
    tables = cursor.fetchall()
    
    # 定义样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    title_font = Font(bold=True, size=14)
    alignment = Alignment(horizontal="center", vertical="center")
    
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 为每个表创建工作表
    for table in tables:
        table_name = table[0]
        ws = wb.create_sheet(title=table_name)
        
        # 获取表结构
        cursor.execute(f"PRAGMA table_info({table_name});")
        columns = cursor.fetchall()
        
        # 获取记录数
        cursor.execute(f"SELECT COUNT(*) FROM {table_name};")
        count = cursor.fetchone()[0]
        
        # 添加表信息
        ws['A1'] = f"表名: {table_name}"
        ws['A1'].font = title_font
        ws['A2'] = f"记录数: {count}"
        ws['A2'].font = Font(bold=True)
        ws['A3'] = f"字段数: {len(columns)}"
        ws['A3'].font = Font(bold=True)
        
        # 添加字段结构信息
        ws['A5'] = "字段结构信息"
        ws['A5'].font = Font(bold=True, size=12)
        
        # 字段结构表头
        structure_headers = ["字段名", "数据类型", "是否非空", "默认值", "主键"]
        for col_idx, header in enumerate(structure_headers, 1):
            cell = ws.cell(row=6, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = alignment
            cell.border = border
        
        # 字段结构数据
        for row_idx, col in enumerate(columns, 7):
            col_id, col_name, col_type, not_null, default_val, pk = col
            data = [
                col_name,
                col_type,
                "是" if not_null else "否",
                str(default_val) if default_val else "",
                "是" if pk else "否"
            ]
            for col_idx, value in enumerate(data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = border
                cell.alignment = alignment
        
        # 添加示例数据
        if count > 0:
            data_start_row = 7 + len(columns) + 2
            ws.cell(row=data_start_row, column=1, value="示例数据 (前5条记录)").font = Font(bold=True, size=12)
            
            # 获取示例数据
            cursor.execute(f"SELECT * FROM {table_name} LIMIT 5;")
            sample_data = cursor.fetchall()
            
            # 列名作为表头
            column_names = [col[1] for col in columns]
            for col_idx, col_name in enumerate(column_names, 1):
                cell = ws.cell(row=data_start_row + 1, column=col_idx, value=col_name)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = alignment
                cell.border = border
            
            # 示例数据
            for row_idx, row_data in enumerate(sample_data, data_start_row + 2):
                for col_idx, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.border = border
                    cell.alignment = alignment
        
        # 调整列宽
        for col_idx in range(1, len(columns) + 1):
            ws.column_dimensions[chr(64 + col_idx)].width = 20
    
    # 创建汇总工作表
    summary_ws = wb.create_sheet(title="数据库结构汇总", index=0)
    
    # 汇总数据
    summary_data = [
        ["数据库文件", db_path],
        ["表数量", len(tables)],
        ["", ""],
        ["表名", "记录数", "字段数", "描述"]
    ]
    
    # 表描述映射
    table_descriptions = {
        "threat_task_performance_nonAI": "威胁任务性能(非AI)",
        "sensor_task_performance_nonAI": "传感器任务性能(非AI)",
        "platform_task_performance_nonAI": "平台任务性能(非AI)",
        "weapon_task_performance_nonAI": "武器任务性能(非AI)",
        "sensor_task_performance_AI": "传感器任务性能(AI)",
        "platform_task_performance_AI": "平台任务性能(AI)",
        "weapon_task_performance_AI": "武器任务性能(AI)",
        "threat_task_performance_AI": "威胁任务性能(AI)"
    }
    
    for table in tables:
        table_name = table[0]
        cursor.execute(f"SELECT COUNT(*) FROM {table_name};")
        count = cursor.fetchone()[0]
        cursor.execute(f"PRAGMA table_info({table_name});")
        columns = cursor.fetchall()
        field_count = len(columns)
        description = table_descriptions.get(table_name, "未知表")
        summary_data.append([table_name, count, field_count, description])
    
    # 写入汇总数据
    for row_idx, row_data in enumerate(summary_data, 1):
        for col_idx, cell_value in enumerate(row_data, 1):
            cell = summary_ws.cell(row=row_idx, column=col_idx, value=cell_value)
            cell.border = border
            
            if row_idx <= 3:  # 前3行是基本信息
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif row_idx == 4:  # 表头行
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = alignment
            else:  # 数据行
                cell.alignment = alignment
    
    # 调整汇总工作表列宽
    summary_ws.column_dimensions['A'].width = 35
    summary_ws.column_dimensions['B'].width = 15
    summary_ws.column_dimensions['C'].width = 10
    summary_ws.column_dimensions['D'].width = 25
    
    conn.close()
    
    # 保存Excel文件
    wb.save(output_file)
    print(f"详细数据库结构已导出到: {output_file}")

if __name__ == "__main__":
    export_detailed_db_to_excel("task_performance_data.db", "task_performance_data_detailed.xlsx")

