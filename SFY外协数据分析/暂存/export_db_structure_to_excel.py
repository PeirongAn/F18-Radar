import sqlite3
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

def export_db_structure_to_excel(db_path, output_file):
    """将数据库结构导出到Excel文件"""
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()
    
    # 创建Excel工作簿
    wb = Workbook()
    
    # 删除默认工作表
    wb.remove(wb.active)
    
    # 获取所有表名
    cursor.execute("SELECT name FROM sqlite_master WHERE type='table';")
    tables = cursor.fetchall()
    
    # 定义样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 为每个表创建工作表
    for table in tables:
        table_name = table[0]
        ws = wb.create_sheet(title=table_name)
        
        # 获取表结构
        cursor.execute(f"PRAGMA table_info({table_name});")
        columns = cursor.fetchall()
        
        # 获取记录数
        cursor.execute(f"SELECT COUNT(*) FROM {table_name};")
        count = cursor.fetchone()[0]
        
        # 准备数据
        data = []
        data.append(["字段名", "数据类型", "是否非空", "默认值", "主键"])
        
        for col in columns:
            col_id, col_name, col_type, not_null, default_val, pk = col
            data.append([
                col_name,
                col_type,
                "是" if not_null else "否",
                str(default_val) if default_val else "",
                "是" if pk else "否"
            ])
        
        # 添加记录数信息
        data.append(["", "", "", "", ""])
        data.append(["记录数", str(count), "", "", ""])
        
        # 写入数据到工作表
        for row_idx, row_data in enumerate(data, 1):
            for col_idx, cell_value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=cell_value)
                cell.border = border
                
                # 设置表头样式
                if row_idx == 1:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_alignment
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # 调整列宽
        column_widths = [25, 15, 10, 15, 8]
        for i, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + i)].width = width
    
    # 创建汇总工作表
    summary_ws = wb.create_sheet(title="数据库结构汇总", index=0)
    
    # 汇总数据
    summary_data = [
        ["数据库文件", db_path],
        ["表数量", len(tables)],
        ["", ""],
        ["表名", "记录数", "字段数"]
    ]
    
    for table in tables:
        table_name = table[0]
        cursor.execute(f"SELECT COUNT(*) FROM {table_name};")
        count = cursor.fetchone()[0]
        cursor.execute(f"PRAGMA table_info({table_name});")
        columns = cursor.fetchall()
        field_count = len(columns)
        summary_data.append([table_name, count, field_count])
    
    # 写入汇总数据
    for row_idx, row_data in enumerate(summary_data, 1):
        for col_idx, cell_value in enumerate(row_data, 1):
            cell = summary_ws.cell(row=row_idx, column=col_idx, value=cell_value)
            cell.border = border
            
            if row_idx <= 3:  # 前3行是基本信息
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif row_idx == 4:  # 表头行
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            else:  # 数据行
                cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 调整汇总工作表列宽
    summary_ws.column_dimensions['A'].width = 30
    summary_ws.column_dimensions['B'].width = 15
    summary_ws.column_dimensions['C'].width = 10
    
    conn.close()
    
    # 保存Excel文件
    wb.save(output_file)
    print(f"数据库结构已导出到: {output_file}")

if __name__ == "__main__":
    export_db_structure_to_excel("task_performance_data.db", "task_performance_data_structure.xlsx")
