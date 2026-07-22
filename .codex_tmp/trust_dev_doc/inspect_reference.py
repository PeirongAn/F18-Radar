import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document


ROOT = Path(r"D:\codes\F18-Radar\.codex_tmp\trust_dev_doc")
DOCX = ROOT / "reference-clean.docx"
doc = Document(DOCX)

paragraphs = []
for index, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip()
    if text:
        paragraphs.append({
            "index": index,
            "style": paragraph.style.name if paragraph.style else "",
            "alignment": str(paragraph.alignment),
            "text": text,
            "runs": [
                {
                    "text": run.text,
                    "bold": run.bold,
                    "italic": run.italic,
                    "size_pt": run.font.size.pt if run.font.size else None,
                    "font": run.font.name,
                }
                for run in paragraph.runs if run.text
            ],
        })

tables = []
for table_index, table in enumerate(doc.tables):
    tables.append({
        "index": table_index,
        "style": table.style.name if table.style else "",
        "rows": [
            [cell.text.strip() for cell in row.cells]
            for row in table.rows
        ],
    })

styles = {}
for style in doc.styles:
    if style.type == 1 and style.name in {"Normal", "Title", "Heading 1", "Heading 2", "Heading 3"}:
        styles[style.name] = {
            "font": style.font.name,
            "size_pt": style.font.size.pt if style.font.size else None,
            "bold": style.font.bold,
            "color": str(style.font.color.rgb) if style.font.color and style.font.color.rgb else None,
        }

payload = {"paragraphs": paragraphs, "tables": tables, "styles": styles}
(ROOT / "reference-content.json").write_text(
    json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
)

images_dir = ROOT / "reference-images"
images_dir.mkdir(exist_ok=True)
with ZipFile(DOCX) as archive:
    for name in archive.namelist():
        if name.startswith("word/media/"):
            (images_dir / Path(name).name).write_bytes(archive.read(name))

print(json.dumps(payload, ensure_ascii=False, indent=2))
