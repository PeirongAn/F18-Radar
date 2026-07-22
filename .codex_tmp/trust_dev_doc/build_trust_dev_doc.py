# -*- coding: utf-8 -*-
from pathlib import Path
import math, re
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(r"D:\codes\F18-Radar"); DIR=ROOT/".codex_tmp"/"trust_dev_doc"
SRC=DIR/"trust_dev_doc.md"; OUT=DIR/"????????-????.docx"
UI=DIR/"ui_aoi_layout.png"; FLOW=DIR/"data_flow.png"
GREEN="0C6B43"; DARK="073B2A"; MID="DDEFE7"; LIGHT="F2F8F5"; BLUE="EAF1F5"; GRAY="67746E"; WHITE="FFFFFF"

def fnt(names,size):
    for n in names:
        p=Path(r"C:\Windows\Fonts")/n
        if p.exists(): return ImageFont.truetype(str(p),size)
    return ImageFont.load_default()
FB=fnt(["msyh.ttc","simsun.ttc"],20); FH=fnt(["msyhbd.ttc","simhei.ttf"],26); FS=fnt(["msyh.ttc"],16)

def center(d,box,text,font,fill):
    x0,y0,x1,y1=box; b=d.multiline_textbbox((0,0),text,font=font,spacing=4,align="center")
    w,h=b[2]-b[0],b[3]-b[1]
    d.multiline_text(((x0+x1-w)/2,(y0+y1-h)/2),text,font=font,fill=fill,spacing=4,align="center")

def diagrams():
    im=Image.new("RGB",(1600,900),"#06140d"); d=ImageDraw.Draw(im)
    d.rectangle((0,0,1600,68),fill="#08281b",outline="#0c6b43",width=2)
    d.text((34,18),"????? / ?? / ???? / ?? / ??",font=FB,fill="#69e6aa")
    d.rounded_rectangle((30,95,1045,855),10,fill="#071a11",outline="#168356",width=3)
    d.text((55,120),"???????",font=FH,fill="#72f2b7")
    d.rounded_rectangle((90,175,945,780),8,fill="#081f15",outline="#6fd59e",width=2)
    d.text((120,205),"RADAR??????????\nSA??????? / ???",font=FB,fill="#b9e8d0")
    d.ellipse((485,375,555,445),outline="#f2cf43",width=4); d.text((462,455),"AI????",font=FS,fill="#f2cf43")
    for box in [(185,595,355,655),(410,610,580,670),(635,580,805,640)]:
        d.rounded_rectangle(box,8,fill="#113d2a",outline="#55c58c",width=2); center(d,box,"????",FS,"#bce9d1")
    d.text((70,805),"AOI?left_ai_target",font=FS,fill="#f4d66d"); d.text((530,805),"AOI?left_candidate_list",font=FS,fill="#78d7a7")
    d.rounded_rectangle((1070,95,1570,855),10,fill="#071a11",outline="#168356",width=3)
    d.text((1092,115),"??????????",font=FH,fill="#72f2b7")
    panels=[(155,285,"AI?????????","right_ai_history_accuracy"),(295,405,"AI????","right_recommendation"),(415,585,"????????","right_candidate_list"),(595,725,"TDC?? / ??????","right_detail"),(735,830,"????????","right_comparison")]
    for y0,y1,label,aoi in panels:
        d.rounded_rectangle((1090,y0,1550,y1),8,fill="#102d22",outline="#2da271",width=2)
        d.text((1110,y0+14),label,font=FB,fill="#d7f2e4"); d.text((1110,y1-32),aoi,font=FS,fill="#77cfa3")
    im.save(UI)
    im=Image.new("RGB",(1600,900),"#f6f8f7"); d=ImageDraw.Draw(im); d.text((50,32),"????????",font=FH,fill="#073b2a")
    boxes=[(45,120,345,300,"??????\nRADAR / SA\nTDC?Button1/2/3","#ddefe7"),(420,120,760,300,"WebSocket??\n???? / ????\n????","#eaf1f5"),(835,120,1190,300,"?????\n??????\n???????","#f4ead9"),(1255,120,1560,300,"?????\n?????\n????","#e9e2f2"),(420,470,760,650,"AOI????\n?? + ???\n+ ????","#eaf1f5"),(835,470,1190,650,"GazeService\n???????\nAOI??","#ddefe7"),(1255,470,1560,650,"????\nraw_gaze.jsonl\n???????","#e9e2f2")]
    for x0,y0,x1,y1,t,c in boxes:
        d.rounded_rectangle((x0,y0,x1,y1),18,fill=c,outline="#0c6b43",width=3); center(d,(x0+8,y0+8,x1-8,y1-8),t,FB,"#20302a")
    for (x0,y0),(x1,y1) in [((345,210),(420,210)),((760,210),(835,210)),((1190,210),(1255,210)),((345,260),(420,540)),((760,560),(835,560)),((1190,560),(1255,560)),((1010,300),(1010,470))]:
        d.line((x0,y0,x1,y1),fill="#0c6b43",width=5); a=math.atan2(y1-y0,x1-x0)
        d.polygon([(x1,y1),(x1-18*math.cos(a-.45),y1-18*math.sin(a-.45)),(x1-18*math.cos(a+.45),y1-18*math.sin(a+.45))],fill="#0c6b43")
    d.rounded_rectangle((45,725,1560,842),16,fill="#073b2a")
    center(d,(65,735,1540,832),"?????task_group_id + task_id / trial_id\n????????????AOI???????",FB,"#effaf4")
    im.save(FLOW)

def font(run,ea="??",latin="Times New Roman",size=10.5,bold=None,color=None):
    run.font.name=latin; run._element.rPr.rFonts.set(qn("w:eastAsia"),ea); run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if color: run.font.color.rgb=RGBColor.from_string(color)
def pf(p,after=3): p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.25
def shade(cell,color):
    pr=cell._tc.get_or_add_tcPr(); n=pr.find(qn("w:shd"))
    if n is None: n=OxmlElement("w:shd"); pr.append(n)
    n.set(qn("w:fill"),color)
def margins(cell):
    pr=cell._tc.get_or_add_tcPr(); mar=OxmlElement("w:tcMar"); pr.append(mar)
    for k,v in [("top",80),("start",90),("bottom",80),("end",90)]:
        n=OxmlElement("w:"+k); n.set(qn("w:w"),str(v)); n.set(qn("w:type"),"dxa"); mar.append(n)
def width(cell,cm):
    pr=cell._tc.get_or_add_tcPr(); n=OxmlElement("w:tcW"); n.set(qn("w:w"),str(int(Cm(cm).twips))); n.set(qn("w:type"),"dxa"); pr.append(n)
def heading(doc,text,level):
    p=doc.add_paragraph(style=f"Heading {level}"); font(p.add_run(text),"??",size={1:15,2:13,3:11.5}[level],bold=True,color=DARK if level<3 else "20302A")
def paragraph(doc,text="",style=None):
    p=doc.add_paragraph(style=style); pf(p); font(p.add_run(text)); return p
def table(doc,rows):
    headers=rows[0]; data=rows[1:]; n=len(headers); ws={2:[4.2,10.1],3:[3.2,7.1,4.0],4:[2.2,3.8,6.3,2.0]}.get(n,[14.3/n]*n)
    t=doc.add_table(rows=1,cols=n); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    tr=t.rows[0]; rp=tr._tr.get_or_add_trPr(); h=OxmlElement("w:tblHeader"); h.set(qn("w:val"),"true"); rp.append(h)
    for i,x in enumerate(headers):
        c=tr.cells[i]; shade(c,DARK); margins(c); width(c,ws[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf(p,0); font(p.add_run(x),"??",size=8.6,bold=True,color=WHITE)
    for ri,row in enumerate(data):
        cells=t.add_row().cells; t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i,x in enumerate(row):
            c=cells[i]; shade(c,WHITE if ri%2==0 else LIGHT); margins(c); width(c,ws[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            p=c.paragraphs[0]; pf(p,0); font(p.add_run(x),size=8.5)
    paragraph(doc,""); return t
def code(doc,text):
    p=doc.add_paragraph(); pf(p,5); p.paragraph_format.left_indent=Cm(.25); p.paragraph_format.right_indent=Cm(.25)
    pr=p._p.get_or_add_pPr(); sh=OxmlElement("w:shd"); sh.set(qn("w:fill"),BLUE); pr.append(sh)
    b=OxmlElement("w:pBdr"); l=OxmlElement("w:left"); l.set(qn("w:val"),"single"); l.set(qn("w:sz"),"16"); l.set(qn("w:color"),GREEN); b.append(l); pr.append(b)
    font(p.add_run(text),"??","Consolas",8.2,color="18352A")
def note(doc,label,text,kind):
    colors={"info":(MID,DARK),"warn":("FFF3DD","C47A12"),"gap":("FBE9E7","A23A35")}; fill,col=colors.get(kind,colors["info"])
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; c=t.cell(0,0); shade(c,fill); margins(c); width(c,14.3)
    p=c.paragraphs[0]; pf(p,0); font(p.add_run(label+"?"),"??",size=10,bold=True,color=col); font(p.add_run(text),size=10); paragraph(doc,"")
def field(p,name):
    r=p.add_run(); a=OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"),"begin"); b=OxmlElement("w:instrText"); b.text=name; c=OxmlElement("w:fldChar"); c.set(qn("w:fldCharType"),"end"); r._r.extend([a,b,c])
def setup(doc):
    s=doc.sections[0]; s.page_width=Cm(21); s.page_height=Cm(29.7); s.left_margin=s.right_margin=Inches(1.25); s.top_margin=s.bottom_margin=Inches(1)
    p=s.header.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(p.add_run("???????? ? ????"),"??",size=8.5,color=GRAY)
    p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run("V1.0 ? "),size=8.5,color=GRAY); field(p,"PAGE"); font(p.add_run(" / "),size=8.5,color=GRAY); field(p,"NUMPAGES")
    st=doc.styles["Normal"]; st.font.name="Times New Roman"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"??"); st.font.size=Pt(10.5)
    for n in ["List Bullet","List Number"]:
        st=doc.styles[n]; st.font.name="Times New Roman"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"??"); st.font.size=Pt(10.5)
def cover(doc):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(90); font(p.add_run("????????"),"??",size=23,bold=True,color=DARK)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(10); font(p.add_run("????"),"??",size=28,bold=True,color=DARK)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(16); font(p.add_run("UI ?? ? ???? ? ???? ? ???? ? ????"),"??",size=12,color=GRAY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(24); font(p.add_run("???????????????"),size=11)
    doc.add_paragraph().paragraph_format.space_before=Pt(55)
    table(doc,[["??","??"],["????","V1.0"],["????","2026?7?21?"],["????","F18-Radar ????????"],["????","??????RADAR_TARGETING?/ ?????SA_THREAT_ASSESSMENT?"],["????","??????????? WebSocket / ???????"],["?????","?????????????????????"]])
    note(doc,"????","????????????????????????????????????????????????????????","info"); doc.add_page_break()
def parse(doc,text):
    lines=text.splitlines(); i=0; fence="\x60\x60\x60"
    while i<len(lines):
        x=lines[i].rstrip()
        if not x: i+=1; continue
        if x=="---PAGE---": doc.add_page_break(); i+=1; continue
        if x.startswith("![") and "](" in x:
            m=re.match(r"!\[(.*?)\]\((.*?)\)",x); doc.add_picture(str(DIR/m.group(2)),width=Cm(14.3))
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run(m.group(1)),size=9,color=GRAY); i+=1; continue
        if x.startswith(":::"):
            kind,label,txt=x.split("|",2); note(doc,label.strip(),txt.strip(),kind.split()[1]); i+=1; continue
        if x.startswith(fence):
            buf=[]; i+=1
            while i<len(lines) and not lines[i].startswith(fence): buf.append(lines[i]); i+=1
            code(doc,"\n".join(buf)); i+=1; continue
        if x.startswith("|"):
            buf=[]
            while i<len(lines) and lines[i].startswith("|"): buf.append(lines[i]); i+=1
            rows=[[c.strip() for c in r.strip("|").split("|")] for r in buf]
            if len(rows)>1 and all(set(c)<=set("-: ") for c in rows[1]): rows.pop(1)
            table(doc,rows); continue
        m=re.match(r"^(#{1,3})\s+(.*)",x)
        if m: heading(doc,m.group(2),len(m.group(1))); i+=1; continue
        if x.startswith("- "): paragraph(doc,x[2:],"List Bullet"); i+=1; continue
        if re.match(r"^\d+\.\s",x): paragraph(doc,re.sub(r"^\d+\.\s","",x),"List Number"); i+=1; continue
        buf=[x]; i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r"^(#{1,3})\s|^- |^\d+\.\s|^\||^\x60\x60\x60|^:::|^!\[|^---PAGE---",lines[i]):
            buf.append(lines[i].strip()); i+=1
        paragraph(doc," ".join(buf))
def main():
    diagrams(); doc=Document(); setup(doc); cover(doc); parse(doc,SRC.read_text(encoding="utf-8"))
    for p in doc.paragraphs:
        pr=p._p.get_or_add_pPr()
        if pr.find(qn("w:widowControl")) is None: pr.append(OxmlElement("w:widowControl"))
    doc.core_properties.title="?????????????"; doc.core_properties.author="F18-Radar ???"; doc.save(OUT); print(OUT)
if __name__=="__main__": main()
