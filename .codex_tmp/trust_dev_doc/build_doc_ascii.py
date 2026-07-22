from pathlib import Path
import math,re
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT,WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm,Inches,Pt,RGBColor
from PIL import Image,ImageDraw,ImageFont

D=Path(r"D:\codes\F18-Radar\.codex_tmp\trust_dev_doc")
SRC=D/"trust_dev_doc_utf8.md"; OUT=D/"trust_control_development_guide.docx"
UI=D/"ui_layout_en.png"; FLOW=D/"data_flow_en.png"
GREEN="0C6B43";DARK="073B2A";MID="DDEFE7";LIGHT="F2F8F5";BLUE="EAF1F5";GRAY="67746E";WHITE="FFFFFF"

def imgfont(size,bold=False):
    names=["arialbd.ttf","arial.ttf"] if bold else ["arial.ttf"]
    for n in names:
        p=Path(r"C:\Windows\Fonts")/n
        if p.exists():return ImageFont.truetype(str(p),size)
    return ImageFont.load_default()
FB=imgfont(20);FH=imgfont(27,True);FS=imgfont(16)
def ctr(d,b,t,f,c):
    x0,y0,x1,y1=b; q=d.multiline_textbbox((0,0),t,font=f,spacing=4,align="center");w=q[2]-q[0];h=q[3]-q[1]
    d.multiline_text(((x0+x1-w)/2,(y0+y1-h)/2),t,font=f,fill=c,spacing=4,align="center")
def diagrams():
    im=Image.new("RGB",(1600,900),"#06140d");d=ImageDraw.Draw(im)
    d.rectangle((0,0,1600,68),fill="#08281b",outline="#0c6b43",width=2);d.text((35,18),"TASK STATUS / USER / TYPE / PROGRESS / DIFFICULTY",font=FB,fill="#69e6aa")
    d.rounded_rectangle((30,95,1045,855),10,fill="#071a11",outline="#168356",width=3);d.text((55,120),"LEFT TASK VIEW",font=FH,fill="#72f2b7")
    d.rounded_rectangle((90,175,945,780),8,fill="#081f15",outline="#6fd59e",width=2);ctr(d,(100,190,930,740),"RADAR: SITUATION + TARGETS\nSA: THREAT LIST + SITUATION\n\nAI RECOMMENDED TARGET\n\nCANDIDATE TARGETS",FB,"#b9e8d0")
    d.text((70,805),"AOI: left_ai_target",font=FS,fill="#f4d66d");d.text((520,805),"AOI: left_candidate_list",font=FS,fill="#78d7a7")
    d.rounded_rectangle((1070,95,1570,855),10,fill="#071a11",outline="#168356",width=3);d.text((1092,115),"RIGHT TRUST PANEL",font=FH,fill="#72f2b7")
    ps=[(155,285,"AI HISTORY ACCURACY","right_ai_history_accuracy"),(295,405,"AI RECOMMENDATION","right_recommendation"),(415,585,"CANDIDATE LIST","right_candidate_list"),(595,725,"TDC / MANUAL REVIEW DETAIL","right_detail"),(735,830,"HUMAN-AI COMPARISON","right_comparison")]
    for y0,y1,l,a in ps:
        d.rounded_rectangle((1090,y0,1550,y1),8,fill="#102d22",outline="#2da271",width=2);d.text((1110,y0+14),l,font=FB,fill="#d7f2e4");d.text((1110,y1-32),a,font=FS,fill="#77cfa3")
    im.save(UI)
    im=Image.new("RGB",(1600,900),"#f6f8f7");d=ImageDraw.Draw(im);d.text((50,32),"TRUST CONTROL DATA LOOP",font=FH,fill="#073b2a")
    bs=[(45,120,345,300,"FRONTEND\nRADAR / SA\nTDC + JOYSTICK","#ddefe7"),(420,120,760,300,"WEBSOCKET\nSTART / EVENTS\nCONFIRM","#eaf1f5"),(835,120,1190,300,"SERVER\nSNAPSHOT CHECK\nTRUST SETTLEMENT","#f4ead9"),(1255,120,1560,300,"TASK DB\nRUNS / EVENTS\nOUTCOMES","#e9e2f2"),(420,470,760,650,"AOI SNAPSHOT\nGEOMETRY + VISIBLE\n+ BINDING","#eaf1f5"),(835,470,1190,650,"GAZE SERVICE\nREVISION + FRAME\nAOI HITS","#ddefe7"),(1255,470,1560,650,"GAZE STORAGE\nRAW JSONL\nSNAPSHOT + STATS","#e9e2f2")]
    for x0,y0,x1,y1,t,c in bs:d.rounded_rectangle((x0,y0,x1,y1),18,fill=c,outline="#0c6b43",width=3);ctr(d,(x0+8,y0+8,x1-8,y1-8),t,FB,"#20302a")
    for (x0,y0),(x1,y1) in [((345,210),(420,210)),((760,210),(835,210)),((1190,210),(1255,210)),((345,260),(420,540)),((760,560),(835,560)),((1190,560),(1255,560)),((1010,300),(1010,470))]:
        d.line((x0,y0,x1,y1),fill="#0c6b43",width=5);a=math.atan2(y1-y0,x1-x0);d.polygon([(x1,y1),(x1-18*math.cos(a-.45),y1-18*math.sin(a-.45)),(x1-18*math.cos(a+.45),y1-18*math.sin(a+.45))],fill="#0c6b43")
    d.rounded_rectangle((45,725,1560,842),16,fill="#073b2a");ctr(d,(65,735,1540,832),"LINK: task_group_id + task_id / trial_id\nBUSINESS TIME: ms; GAZE AND AOI INTERVAL: us",FB,"#effaf4");im.save(FLOW)

def rfont(r,ea="SimSun",latin="Times New Roman",size=10.5,bold=None,color=None):
    r.font.name=latin;r._element.rPr.rFonts.set(qn("w:eastAsia"),ea);r.font.size=Pt(size)
    if bold is not None:r.bold=bold
    if color:r.font.color.rgb=RGBColor.from_string(color)
def pformat(p,a=3):p.paragraph_format.space_after=Pt(a);p.paragraph_format.line_spacing=1.25
def shade(c,x):
    pr=c._tc.get_or_add_tcPr();n=pr.find(qn("w:shd"))
    if n is None:n=OxmlElement("w:shd");pr.append(n)
    n.set(qn("w:fill"),x)
def cmarg(c):
    pr=c._tc.get_or_add_tcPr();m=OxmlElement("w:tcMar");pr.append(m)
    for k,v in [("top",80),("start",90),("bottom",80),("end",90)]:
        n=OxmlElement("w:"+k);n.set(qn("w:w"),str(v));n.set(qn("w:type"),"dxa");m.append(n)
def cwidth(c,v):
    n=OxmlElement("w:tcW");n.set(qn("w:w"),str(int(Cm(v).twips)));n.set(qn("w:type"),"dxa");c._tc.get_or_add_tcPr().append(n)
def para(doc,t="",style=None):
    p=doc.add_paragraph(style=style);pformat(p);rfont(p.add_run(t));return p
def heading(doc,t,l):
    p=doc.add_paragraph(style=f"Heading {l}");rfont(p.add_run(t),"SimHei",size={1:15,2:13,3:11.5}[l],bold=True,color=DARK if l<3 else "20302A")
def tbl(doc,rows):
    n=len(rows[0]);ws={2:[4.2,10.1],3:[3.2,7.1,4.0],4:[2.2,3.8,6.3,2.0],5:[2.5,2.7,2.8,3.2,3.1]}.get(n,[14.3/n]*n)
    t=doc.add_table(rows=1,cols=n);t.style="Table Grid";t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    rp=t.rows[0]._tr.get_or_add_trPr();h=OxmlElement("w:tblHeader");h.set(qn("w:val"),"true");rp.append(h)
    for i,x in enumerate(rows[0]):
        c=t.rows[0].cells[i];shade(c,DARK);cmarg(c);cwidth(c,ws[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pformat(p,0);rfont(p.add_run(x),"SimHei",size=8.5,bold=True,color=WHITE)
    for ri,row in enumerate(rows[1:]):
        cs=t.add_row().cells;t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i,x in enumerate(row):
            c=cs[i];shade(c,WHITE if ri%2==0 else LIGHT);cmarg(c);cwidth(c,ws[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP;p=c.paragraphs[0];pformat(p,0);rfont(p.add_run(x),size=8.4)
    para(doc,"")
def block(doc,t):
    p=doc.add_paragraph();pformat(p,5);p.paragraph_format.left_indent=Cm(.25);p.paragraph_format.right_indent=Cm(.25);pr=p._p.get_or_add_pPr();s=OxmlElement("w:shd");s.set(qn("w:fill"),BLUE);pr.append(s);rfont(p.add_run(t),"Microsoft YaHei","Consolas",8.1,color="18352A")
def note(doc,label,t,k):
    colors={"info":(MID,DARK),"warn":("FFF3DD","C47A12"),"gap":("FBE9E7","A23A35")};f,c=colors.get(k,colors["info"]);tb=doc.add_table(rows=1,cols=1);tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    cell=tb.cell(0,0);shade(cell,f);cmarg(cell);cwidth(cell,14.3);p=cell.paragraphs[0];pformat(p,0);rfont(p.add_run(label+": "),"SimHei",size=10,bold=True,color=c);rfont(p.add_run(t),size=10);para(doc,"")
def field(p,n):
    r=p.add_run();a=OxmlElement("w:fldChar");a.set(qn("w:fldCharType"),"begin");b=OxmlElement("w:instrText");b.text=n;c=OxmlElement("w:fldChar");c.set(qn("w:fldCharType"),"end");r._r.extend([a,b,c])
def setup(doc,title):
    s=doc.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.left_margin=s.right_margin=Inches(1.25);s.top_margin=s.bottom_margin=Inches(1)
    p=s.header.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;rfont(p.add_run(title),"Microsoft YaHei",size=8.5,color=GRAY)
    p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;rfont(p.add_run("V1.0 - "),size=8.5,color=GRAY);field(p,"PAGE");rfont(p.add_run(" / "),size=8.5,color=GRAY);field(p,"NUMPAGES")
    st=doc.styles["Normal"];st.font.name="Times New Roman";st._element.rPr.rFonts.set(qn("w:eastAsia"),"SimSun");st.font.size=Pt(10.5)
def parse(doc,text):
    lines=text.splitlines();meta={};i=0
    while i<len(lines) and lines[i].startswith("@"):
        k,v=lines[i][1:].split(":",1);meta[k]=v;i+=1
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(85);rfont(p.add_run(meta["title"]),"SimHei",size=24,bold=True,color=DARK)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(10);rfont(p.add_run(meta["subtitle"]),"SimHei",size=28,bold=True,color=DARK)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(16);rfont(p.add_run(meta["tagline"]),"Microsoft YaHei",size=12,color=GRAY)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(24);rfont(p.add_run(meta["source"]),size=11);doc.add_page_break()
    fence="\x60\x60\x60"
    while i<len(lines):
        x=lines[i].rstrip()
        if not x:i+=1;continue
        if x=="---PAGE---":doc.add_page_break();i+=1;continue
        if x.startswith("!["):
            m=re.match(r"!\[(.*?)\]\((.*?)\)",x);doc.add_picture(str(D/m.group(2)),width=Cm(14.3));p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;rfont(p.add_run(m.group(1)),size=9,color=GRAY);i+=1;continue
        if x.startswith(":::"):
            k,l,t=x.split("|",2);note(doc,l.strip(),t.strip(),k.split()[1]);i+=1;continue
        if x.startswith(fence):
            b=[];i+=1
            while i<len(lines) and not lines[i].startswith(fence):b.append(lines[i]);i+=1
            block(doc,"\n".join(b));i+=1;continue
        if x.startswith("|"):
            b=[]
            while i<len(lines) and lines[i].startswith("|"):b.append(lines[i]);i+=1
            rows=[[c.strip() for c in z.strip("|").split("|")] for z in b]
            if len(rows)>1 and all(set(c)<=set("-: ") for c in rows[1]):rows.pop(1)
            tbl(doc,rows);continue
        m=re.match(r"^(#{1,3})\s+(.*)",x)
        if m:heading(doc,m.group(2),len(m.group(1)));i+=1;continue
        if x.startswith("- "):para(doc,x[2:],"List Bullet");i+=1;continue
        if re.match(r"^\d+\.\s",x):para(doc,x);i+=1;continue
        b=[x];i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r"^(#{1,3})\s|^- |^\d+\.\s|^\||^\x60\x60\x60|^:::|^!\[|^---PAGE---",lines[i]):b.append(lines[i].strip());i+=1
        para(doc," ".join(b))
    return meta
def main():
    diagrams();text=SRC.read_text(encoding="utf-8");title=text.splitlines()[0].split(":",1)[1];doc=Document();setup(doc,title);meta=parse(doc,text);doc.core_properties.title=meta["title"]+" - "+meta["subtitle"];doc.save(OUT);print(OUT)
if __name__=="__main__":main()
