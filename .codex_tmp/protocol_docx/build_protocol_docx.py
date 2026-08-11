from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"D:\codes\F18-Radar\server\docs\external_task_behavior_pose_protocol.md")
OUTPUT = Path(r"D:\codes\F18-Radar\server\docs\external_task_behavior_pose_protocol_v1.0.docx")

BODY_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"
INK = "172033"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
BORDER = "B8C4D1"
WHITE = "FFFFFF"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name=BODY_FONT, size=None, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr()
    fonts = run._element.rPr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run._element.rPr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_border(paragraph, side="bottom", color=BORDER, size=8, space=4):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    edge = p_bdr.find(qn(f"w:{side}"))
    if edge is None:
        edge = OxmlElement(f"w:{side}")
        p_bdr.append(edge)
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), str(space))
    edge.set(qn("w:color"), color)


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    assert sum(widths) == CONTENT_WIDTH_DXA
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        prevent_row_split(row)
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            cell.width = Inches(widths[idx] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "5")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), BORDER)


def set_code_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for side in ("top", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        edge.set(qn("w:val"), "nil")
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), BLUE)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])
    set_run_font(run, size=8.5, color=MUTED)


def populate_header(header):
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    left = p.add_run("F18-RADAR · 外部任务接口规范")
    set_run_font(left, size=8.5, color=MUTED, bold=True)
    right = p.add_run("\tSCHEMA 1.0")
    set_run_font(right, size=8.5, color=MUTED)
    tabs = p.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(6.5), alignment=2)


def populate_footer(footer):
    fp = footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    prefix = fp.add_run("第 ")
    set_run_font(prefix, size=8.5, color=MUTED)
    add_page_field(fp)
    suffix = fp.add_run(" 页")
    set_run_font(suffix, size=8.5, color=MUTED)


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    section.different_first_page_header_footer = False
    section.header.paragraphs[0].text = ""
    section.even_page_header.paragraphs[0].text = ""
    populate_footer(section.footer)
    populate_footer(section.even_page_footer)


def configure_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.25
    pf.widow_control = True

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = styles[style_name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        spf = style.paragraph_format
        spf.space_before = Pt(before)
        spf.space_after = Pt(after)
        spf.line_spacing = 1.0
        spf.left_indent = Inches(0)
        spf.right_indent = Inches(0)
        spf.first_line_indent = Inches(0)
        spf.keep_with_next = True
        spf.keep_together = True
        spf.page_break_before = False

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = BODY_FONT
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string(INK)
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        lpf = style.paragraph_format
        lpf.left_indent = Inches(0.375)
        lpf.first_line_indent = Inches(-0.188)
        lpf.space_after = Pt(4)
        lpf.line_spacing = 1.25


INLINE_TOKEN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")


def add_inline(paragraph, text, size=11, color=INK, bold=False):
    cursor = 0
    for match in INLINE_TOKEN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, color=color, bold=bold)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name=CODE_FONT, size=max(8.5, size - 0.7), color=DARK_BLUE)
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, color=color, bold=bold)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("接口规范")
    set_run_font(r, size=10, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("外部任务行为与姿态数据接口说明")
    set_run_font(r, size=24, color=INK, bold=True)
    r.add_break()
    r = p.add_run("（Web 对接版）")
    set_run_font(r, size=20, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("适用于 PLATFORM_CONTROL 与 WEAPON_FIRING 外部任务")
    set_run_font(r, size=12.5, color=MUTED)
    set_paragraph_border(p, side="bottom", color=BLUE, size=10, space=8)

    metadata = [
        ("协议版本", "1.0"),
        ("传输方式", "WebSocket"),
        ("发送策略", "状态变化触发：开始 *_start，结束 *_end"),
        ("任务关联", "服务端解析 task_runs.task_id"),
        ("时间基准", "Unix 毫秒时间戳"),
        ("修订日期", "2026-07-29"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(f"{label}：")
        set_run_font(r, size=9.5, color=MUTED, bold=True)
        r = p.add_run(value)
        set_run_font(r, size=9.5, color=INK)


def choose_widths(rows):
    cols = len(rows[0])
    if cols == 2:
        first_values = [len(r[0]) for r in rows]
        second_values = [len(r[1]) for r in rows]
        if max(first_values) <= 18 and max(second_values) > 35:
            return [2700, 6660]
        return [3900, 5460]
    if cols == 3:
        return [2400, 2200, 4760]
    if cols == 4:
        return [2200, 1550, 1300, 4310]
    return [CONTENT_WIDTH_DXA // cols] * (cols - 1) + [
        CONTENT_WIDTH_DXA - (CONTENT_WIDTH_DXA // cols) * (cols - 1)
    ]


def add_markdown_table(doc, rows):
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    widths = choose_widths(rows)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_table_borders(table)
    for ridx, row in enumerate(rows):
        for cidx, value in enumerate(row):
            cell = table.cell(ridx, cidx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if ridx < len(rows) - 1:
                p.paragraph_format.keep_with_next = True
            add_inline(p, value, size=9.2, color=INK, bold=(ridx == 0))
            if ridx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
        if ridx == 0:
            set_repeat_table_header(table.rows[ridx])
    return table


def add_code_block(doc, lines, language):
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_code_table_borders(table)
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_margins(cell, top=60, start=120, bottom=60, end=120)

    entries = []
    if language:
        entries.append((language.upper(), 7.5, MUTED, True))
    entries.extend((line, 8.3, INK, False) for line in lines)

    first = cell.paragraphs[0]
    for idx, (text, size, color, bold) in enumerate(entries):
        p = first if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_with_next = idx < len(entries) - 1
        run = p.add_run(text if text else " ")
        set_run_font(run, name=CODE_FONT, size=size, color=color, bold=bold)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(1)
    spacer_run = spacer.add_run(" ")
    set_run_font(spacer_run, size=1, color=WHITE)


def add_quote(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.2
    set_paragraph_shading(p, CALLOUT)
    set_paragraph_border(p, side="left", color=BLUE, size=14, space=6)
    add_inline(p, text, size=10.5, color=DARK_BLUE)


def is_table_separator(line):
    parts = [part.strip() for part in line.strip().strip("|").split("|")]
    return bool(parts) and all(re.fullmatch(r":?-{3,}:?", part) for part in parts)


def parse_markdown(doc, text):
    lines = text.splitlines()
    i = 0
    first_h1_seen = False
    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            language = stripped[3:].strip()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code_block(doc, code_lines, language)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            rows = [[part.strip() for part in stripped.strip("|").split("|")]]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([part.strip() for part in lines[i].strip().strip("|").split("|")])
                i += 1
            add_markdown_table(doc, rows)
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2)
            if level == 1 and not first_h1_seen:
                first_h1_seen = True
                i += 1
                continue
            mapped_level = min(max(level - 1, 1), 3)
            p = doc.add_paragraph(style=f"Heading {mapped_level}")
            p.paragraph_format.left_indent = Inches(0)
            p.paragraph_format.right_indent = Inches(0)
            p.paragraph_format.first_line_indent = Inches(0)
            add_inline(
                p,
                title,
                size={1: 16, 2: 13, 3: 12}[mapped_level],
                color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[mapped_level],
                bold=True,
            )
            i += 1
            continue

        if stripped.startswith(">"):
            add_quote(doc, stripped[1:].strip())
            i += 1
            continue

        bullet = re.match(r"^-\s+(.*)$", stripped)
        if bullet:
            value = bullet.group(1)
            if value.startswith("[ ] "):
                value = "☐ " + value[4:]
            elif value.startswith("[x] ") or value.startswith("[X] "):
                value = "☑ " + value[4:]
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, value)
            i += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, numbered.group(1))
            i += 1
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (
                not nxt
                or nxt.startswith("#")
                or nxt.startswith("```")
                or nxt.startswith("|")
                or nxt.startswith(">")
                or re.match(r"^-\s+", nxt)
                or re.match(r"^\d+\.\s+", nxt)
            ):
                break
            paragraph_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(paragraph_lines))


def add_document_properties(doc):
    props = doc.core_properties
    props.title = "外部任务行为与姿态数据接口说明（Web 对接版）"
    props.subject = "PLATFORM_CONTROL 与 WEAPON_FIRING 外部任务行为及姿态角协议"
    props.keywords = "F18-Radar, WebSocket, behavior, pose, Pitch, Yaw, Roll, task_runs"
    props.comments = "Converted from external_task_behavior_pose_protocol.md"


def main():
    source_text = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    configure_section(section)
    configure_styles(doc)
    add_document_properties(doc)
    add_title_block(doc)
    parse_markdown(doc, source_text)

    for paragraph in doc.paragraphs:
        if paragraph.style.name == "Normal":
            paragraph.paragraph_format.widow_control = True

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
